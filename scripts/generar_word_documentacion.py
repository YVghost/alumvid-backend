import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

doc = docx.Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Colors
COLOR_PRIMARY = RGBColor(30, 58, 138)     # Navy #1E3A8A
COLOR_SECONDARY = RGBColor(14, 116, 144) # Cyan/Blue #0E7490
COLOR_DARK = RGBColor(15, 23, 42)        # Slate 900
COLOR_GRAY = RGBColor(100, 116, 139)     # Slate 500

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

# Title
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(4)
run_title = title_p.add_run('Documentación Técnica: Optimización de Rutas Logísticas en AlumVid')
run_title.font.name = 'Calibri'
run_title.font.size = Pt(20)
run_title.font.bold = True
run_title.font.color.rgb = COLOR_PRIMARY

# Subtitle
sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_after = Pt(14)
run_sub = sub_p.add_run('De la Teoría del Vehicle Routing Problem (VRP) a la Heurística Voraz (Greedy Search) con Google Routes API y Re-enrutado Dinámico (VRSP)')
run_sub.font.name = 'Calibri'
run_sub.font.size = Pt(11.5)
run_sub.font.italic = True
run_sub.font.color.rgb = COLOR_SECONDARY

# Metadata Box
meta_data = [
    ('Institución:', 'Universidad de Las Américas (UDLA) - Facultad de Ingeniería y Ciencias Aplicadas'),
    ('Proyecto de Titulación:', 'Plataforma de control y monitoreo en tiempo real para la optimización de inventario muerto y rutas logísticas en una microempresa de vidrios (AlumVid)'),
    ('Estudiantes:', 'Anthonny Jossue Mosquera Rivadeneira, Mateo Michael Ortega Vacas'),
    ('Tutor Guía:', 'Prof. Mario Salvador González Rodríguez'),
    ('Bodega Central:', 'Arq-AlumVid (Coordenadas: -0.089290, -78.439607) – Calderón / Llano Grande, Quito Norte'),
    ('Fecha / Ubicación:', 'Octubre 2026 – Quito, Ecuador')
]

meta_table = doc.add_table(rows=len(meta_data), cols=2)
meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for idx, (label, val) in enumerate(meta_data):
    row = meta_table.rows[idx]
    c1, c2 = row.cells[0], row.cells[1]
    c1.width = Inches(1.8)
    c2.width = Inches(4.7)
    set_cell_background(c1, 'F1F5F9')
    set_cell_background(c2, 'F8FAFC')
    set_cell_margins(c1, 80, 80, 100, 100)
    set_cell_margins(c2, 80, 80, 100, 100)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run(label)
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = COLOR_PRIMARY
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(val)
    r2.font.size = Pt(9.5)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

md_file = os.path.join('DOCUMENTACION', 'DOCUMENTACION_OPTIMIZACION_RUTAS_VRP_GREEDY.md')
with open(md_file, 'r', encoding='utf-8') as f:
    lines = f.readlines()

in_table = False
table_lines = []

def flush_table(t_lines):
    if not t_lines: return
    headers = [c.strip() for c in t_lines[0].split('|')[1:-1]]
    rows = []
    for line in t_lines[2:]:
        if '|' in line:
            rows.append([c.strip() for c in line.split('|')[1:-1]])
    if not headers or not rows: return
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = t.cell(0, j)
        set_cell_background(cell, '1E3A8A')
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
    for i, r_data in enumerate(rows):
        bg = 'F8FAFC' if i % 2 == 0 else 'FFFFFF'
        for j, val in enumerate(r_data):
            cell = t.cell(i + 1, j)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 120, 120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            # Clean formatting in table cell
            cleaned_val = val.replace('$O(N^2)$', 'O(N²)').replace('$\\rightarrow$', '→')
            parts = cleaned_val.split('**')
            for p_idx, part in enumerate(parts):
                r = p.add_run(part)
                r.font.size = Pt(9)
                if p_idx % 2 == 1:
                    r.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

started_content = False

for line in lines:
    raw = line.rstrip('\r\n')
    
    if not started_content:
        if raw.startswith('## 1. Contexto'):
            started_content = True
        else:
            continue

    # Clean text of raw math delimiters
    clean_line = raw.replace('$O(N^2)$', 'O(N²)').replace('$\\rightarrow$', '→').replace('$$', '')

    if clean_line.startswith('|'):
        in_table = True
        table_lines.append(clean_line)
        continue
    elif in_table:
        in_table = False
        flush_table(table_lines)
        table_lines = []

    if clean_line.startswith('## '):
        h2 = doc.add_heading(clean_line[3:], level=1)
        h2.paragraph_format.space_before = Pt(16)
        h2.paragraph_format.space_after = Pt(6)
        for r in h2.runs:
            r.font.color.rgb = COLOR_PRIMARY
            r.font.size = Pt(13.5)
    elif clean_line.startswith('### '):
        h3 = doc.add_heading(clean_line[4:], level=2)
        h3.paragraph_format.space_before = Pt(12)
        h3.paragraph_format.space_after = Pt(4)
        for r in h3.runs:
            r.font.color.rgb = COLOR_SECONDARY
            r.font.size = Pt(11.5)
    elif clean_line.startswith('#### '):
        h4 = doc.add_heading(clean_line[5:], level=3)
        h4.paragraph_format.space_before = Pt(8)
        h4.paragraph_format.space_after = Pt(3)
        for r in h4.runs:
            r.font.color.rgb = COLOR_DARK
            r.font.size = Pt(10.5)
            r.font.bold = True
    elif clean_line.strip().startswith('* ') or clean_line.strip().startswith('- '):
        stripped = clean_line.strip()
        leading_spaces = len(clean_line) - len(clean_line.lstrip(' '))
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        if leading_spaces >= 2:
            bp.paragraph_format.left_indent = Inches(0.4)
        text = stripped[2:]
        parts = text.split('**')
        for idx, part in enumerate(parts):
            r = bp.add_run(part)
            r.font.size = Pt(10)
            if idx % 2 == 1:
                r.font.bold = True
    elif any(clean_line.strip().startswith(f'{n}. ') for n in range(1, 20)):
        stripped = clean_line.strip()
        dot_idx = stripped.find('. ')
        leading_spaces = len(clean_line) - len(clean_line.lstrip(' '))
        np_par = doc.add_paragraph(style='List Number')
        np_par.paragraph_format.space_after = Pt(3)
        if leading_spaces >= 2:
            np_par.paragraph_format.left_indent = Inches(0.4)
        text = stripped[dot_idx + 2:]
        parts = text.split('**')
        for idx, part in enumerate(parts):
            r = np_par.add_run(part)
            r.font.size = Pt(10)
            if idx % 2 == 1:
                r.font.bold = True
    elif clean_line.startswith('> '):
        qp = doc.add_paragraph()
        qp.paragraph_format.left_indent = Inches(0.4)
        qp.paragraph_format.space_before = Pt(4)
        qp.paragraph_format.space_after = Pt(6)
        text = clean_line[2:]
        r = qp.add_run(text)
        r.font.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_DARK
    elif clean_line.strip() == '---':
        continue
    elif clean_line.startswith('```'):
        continue
    elif clean_line.strip():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        text = clean_line
        parts = text.split('**')
        for idx, part in enumerate(parts):
            r = p.add_run(part)
            r.font.size = Pt(10)
            if idx % 2 == 1:
                r.font.bold = True

if in_table and table_lines:
    flush_table(table_lines)

# Save document
saved_paths = []
output_path = os.path.join('DOCUMENTACION', 'DOCUMENTACION_OPTIMIZACION_RUTAS_VRP_GREEDY.docx')
try:
    doc.save(output_path)
    saved_paths.append(output_path)
    print('Documento principal guardado con exito en:', output_path)
except PermissionError:
    print('Aviso: DOCUMENTACION_OPTIMIZACION_RUTAS_VRP_GREEDY.docx esta abierto en Word.')

# Also save a final copy
alt_path = os.path.join('DOCUMENTACION', 'DOCUMENTACION_OPTIMIZACION_RUTAS_VRP_GREEDY_FINAL.docx')
try:
    doc.save(alt_path)
    saved_paths.append(alt_path)
    print('Version final guardada con exito en:', alt_path)
except PermissionError:
    print('Aviso: DOCUMENTACION_OPTIMIZACION_RUTAS_VRP_GREEDY_FINAL.docx esta abierto en Word.')

print('Guardado completado en:', saved_paths)
